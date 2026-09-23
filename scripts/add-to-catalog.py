# -*- coding: utf-8 -*-
"""Дописывает в Price/Каталог.xlsx товары из свежего прайса, которых там ещё нет.

Зачем: каталог — источник человеческих имён для сайта, но новинки в него попадают
руками. 23.09.2026 их пришло 58 разом (возврат MATRIX), и вписывать столько вручную —
полдня. Скрипт кладёт каждую позицию В СВОЙ РАЗДЕЛ рядом с роднёй и по алфавиту.

Имя приводится к манере каталога: объём слитно («90 мл.» -> «90мл»), лишние пробелы
схлопываются, очевидные сокращения раскрываются («св шатен» -> «светлый шатен»),
артикул остаётся последним словом — на нём держится вся сверка.

LEBEL пропускается: для него каталог отключён как источник (решение 22.09.2026),
имена даёт бланк поставщика.

ВАЖНО, как устроена запись. Разделы пересобираются ЦЕЛИКОМ, а не вставляются строки
по одной. Вставка по одной ломается дважды, и оба раза тихо: номера строк уезжают от
предыдущих вставок (запомненная строка 839 перестаёт указывать туда, куда указывала,
и товар попадает в соседний раздел), а порядок внутри раздела выходит случайным —
«5BC, 5MG, 5M, 5BV» вместо алфавита. Поймано на первом же прогоне 23.09.2026.

Запуск:
    .venv/Scripts/python.exe scripts/add-to-catalog.py            # предпросмотр
    .venv/Scripts/python.exe scripts/add-to-catalog.py --apply    # записать

Перед записью кладётся копия Каталог-до-правки-<дата>.xlsx рядом с оригиналом.
Каталог должен быть ЗАКРЫТ в Excel, иначе запись не пройдёт.
"""
import openpyxl, re, sys, shutil, datetime
CAT = r'Price\Каталог.xlsx'
PRICE = r'public/prices/price-current.xlsx'
APPLY = '--apply' in sys.argv
BRANDS = {'CAREPROST','CONCEPT','DIKSON','ISKARTES','KAPOUS',"L'OREAL",'LEBEL','LEVISSIME',
          'LONDA','MATRIX','OLLIN','SCHWARZKOPF','WELLA','Корея','ОДНОРАЗОВАЯ ПРОДУКЦИЯ'}
# Родню по имени не найти — указываем раздел по ИМЕНИ заголовка, а не по номеру строки.
MANUAL_SECTION = {'2646468': ('SCHWARZKOPF', 'Средства для укладки волос - SILHOUETTE/Professionnelle')}

def art(n):
    t = str(n).strip().split()[-1] if str(n).strip() else ''
    return t.rstrip('.').lower()

def tidy(name):
    s = re.sub(r'\s+', ' ', str(name)).strip()
    s = re.sub(r'(\d)\s*(мл|г|гр|кг|л)\s*\.?(?=\s|$)', r'\1\2', s)
    s = re.sub(r'\bсв\b(?=\s)', 'светлый', s)
    s = re.sub(r'\bтёмн\b(?=\s)', 'тёмный', s)
    return re.sub(r'\s+', ' ', s).strip()

def is_header(v):
    return bool(v) and not re.search(r'\d{3,}', v) and len(v.split()) <= 8 and art(v) != v.split()[-1].lower().rstrip('.') if v else False

wb = openpyxl.load_workbook(CAT)
ws = wb[wb.sheetnames[0]]
n = ws.max_row
vals = [(str(ws.cell(row=i, column=2).value).strip() if ws.cell(row=i, column=2).value else '')
        for i in range(1, n + 1)]

# разметка: строка -> (бренд, заголовок раздела); и границы разделов
brand_at, sec_at = [''] * n, [''] * n
cur_b = cur_s = ''
sec_start = {}
for i, v in enumerate(vals):
    if v in BRANDS:
        cur_b, cur_s = v, ''
    elif v and not re.search(r'\d', v.split()[-1]):
        cur_s = v
        sec_start[(cur_b, cur_s)] = i
    brand_at[i], sec_at[i] = cur_b, cur_s
cat_arts = {art(v) for v in vals if v}

wp = openpyxl.load_workbook(PRICE, read_only=True, data_only=True)
brand, todo = '', []
for r in wp.active.iter_rows(values_only=True):
    nm = str(r[1]).strip() if len(r) > 1 and r[1] else ''
    pr = r[2] if len(r) > 2 else None
    if not nm: continue
    if nm in BRANDS: brand = nm; continue
    if not isinstance(pr, (int, float)) or brand == 'LEBEL': continue
    if art(nm) not in cat_arts: todo.append((brand, nm))
wp.close()

def family(s):
    w = s.split()
    for k in (4, 3, 2, 1):
        if len(w) >= k: yield ' '.join(w[:k])

add_to, unplaced = {}, []
for b, nm in todo:
    a = art(nm)
    if a in MANUAL_SECTION:
        key = MANUAL_SECTION[a]
    else:
        key = None
        for pref in family(nm):
            for i, v in enumerate(vals):
                if v.startswith(pref) and brand_at[i] == b and sec_at[i]:
                    key = (b, sec_at[i]); break
            if key: break
    if not key: unplaced.append((b, nm)); continue
    add_to.setdefault(key, []).append(tidy(nm))

print(f'разделов затронуто: {len(add_to)} | позиций: {sum(len(v) for v in add_to.values())} | без места: {len(unplaced)}')
for (b, s), items in sorted(add_to.items()):
    print(f'  {b:12} | {s[:44]:46} +{len(items)}')
for b, nm in unplaced: print(f'  БЕЗ МЕСТА: {b} | {nm[:56]}')
if not APPLY:
    print('\nпредпросмотр; для записи — --apply'); sys.exit(0)

shutil.copy2(CAT, CAT.replace('.xlsx', f'-до-правки-{datetime.date.today()}.xlsx'))
# пересобираем разделы СНИЗУ ВВЕРХ, чтобы номера строк выше не уезжали
for (b, s), items in sorted(add_to.items(), key=lambda kv: -sec_start[kv[0]]):
    start = sec_start[(b, s)]                       # индекс заголовка (0-based)
    end = start + 1
    while end < len(vals) and vals[end] and sec_at[end] == s and brand_at[end] == b and vals[end] not in BRANDS:
        end += 1
    block = [vals[i] for i in range(start + 1, end)] + items
    block.sort(key=lambda x: x.lower())
    for off in range(end - start - 1, len(block)):  # дописываем недостающие строки
        ws.insert_rows(start + 2 + off)
    for off, name in enumerate(block):
        ws.cell(row=start + 2 + off, column=2, value=name)
wb.save(CAT)
print(f'\nзаписано {sum(len(v) for v in add_to.values())} позиций в {len(add_to)} разделов')
