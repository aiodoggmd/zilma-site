# -*- coding: utf-8 -*-
"""
Предлагает категорию для новых товаров прайса — по уже размеченной родне.

Зачем: при каждом обновлении в прайсе появляется несколько новых позиций без категории. Без
категории товар не находится через «По категории» в каталоге — только через бренд. Раньше
это проставлялось руками и один раз молча проехало мимо: 59 позиций жили без категории
почти две недели (10.09.2026).

Как считает: товар сравнивается с уже размеченными позициями ТОГО ЖЕ бренда по первым двум
словам названия (оттенок, объём и артикул отброшены). Категория предлагается, только если у
всей найденной родни она ОДНА — иначе показываются варианты и решает человек.

**Два слова, а не три.** Первая версия брала три и ломалась на красителях: третьим словом
там идёт название оттенка («краска колестон перламутровый»), и родня не находилась вовсе.

Скрипт НИЧЕГО не пишет — только предлагает. Разметка проставляется в
src/data/price-categories.json после решения человека.

Запуск: .venv/Scripts/python.exe scripts/suggest-categories.py
"""
import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PRICE = ROOT / 'public/prices/price-current.xlsx'
JOURNAL = ROOT / 'src/data/price-categories.json'

# Названия категорий читаются прямо из xlsx-to-price-items.py — единственного места, где
# они заданы. Дублировать список здесь нельзя: он тихо разъедется с настоящим.
def category_names() -> dict:
    src = (ROOT / 'scripts/xlsx-to-price-items.py').read_text(encoding='utf-8')
    block = re.search(r'CATEGORY_NAMES = \{(.*?)\n\}', src, re.S)
    return dict(re.findall(r'(\d+):\s*"([^"]+)"', block.group(1))) if block else {}


CATEGORY_NAMES = category_names()


def key_of(name: str) -> str:
    """Первые два слова без цифр, кодов оттенков и знаков — «родственный» ключ товара."""
    words = [w for w in re.split(r'[\s,./]+', name.lower())
             if w and not re.fullmatch(r'[\d/\-.,]+', w)]
    return ' '.join(words[:2])


def main() -> None:
    journal = json.loads(JOURNAL.read_text(encoding='utf-8'))
    wb = openpyxl.load_workbook(PRICE, data_only=True)
    ws = wb.active

    # Раскладка листа: B — наименование, F — категория (служебная, скрытая колонка).
    known = defaultdict(set)       # (бренд, ключ) -> {категории}
    unmarked = []                  # (строка, бренд, имя)
    brand = ''
    for row in range(1, ws.max_row + 1):
        name = ws.cell(row, 2).value
        if not name:
            continue
        name = str(name).strip()
        price = ws.cell(row, 3).value
        unit = ws.cell(row, 4).value
        if price is None and unit is None:          # заголовок бренда/линейки
            if name.isupper() or name.upper() == name:
                brand = name
            continue
        cat = str(ws.cell(row, 6).value or '').strip()
        if cat.isdigit():
            known[(brand, key_of(name))].add(int(cat))
        elif not cat:
            unmarked.append((row, brand, name))
        # Иначе это шапка служебной колонки («категории») — не товар и не пропуск.

    if not unmarked:
        print('Все товары размечены — делать нечего.')
        return

    print(f'Без категории: {len(unmarked)}\n')
    for row, br, name in unmarked:
        cats = known.get((br, key_of(name)), set())
        if len(cats) == 1:
            num = next(iter(cats))
            label = CATEGORY_NAMES.get(str(num), '')
            print(f'  строка {row}  {name}')
            print(f'      → {num} {label}  (по родне: {br} / «{key_of(name)}»)')
        elif cats:
            print(f'  строка {row}  {name}')
            print(f'      ? родня размечена по-разному: {sorted(cats)} — решает человек')
        else:
            print(f'  строка {row}  {name}')
            print(f'      ? родни нет ({br} / «{key_of(name)}») — решает человек')
    print(f'\nЖурнал разметки: {JOURNAL.relative_to(ROOT)} ({len(journal)} записей)')


if __name__ == '__main__':
    main()
