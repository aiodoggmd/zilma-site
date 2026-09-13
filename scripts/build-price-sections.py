# -*- coding: utf-8 -*-
"""
Собирает журнал разделов каталога: «имя товара -> раздел внутри бренда».

ЗАЧЕМ. В прайсе из 1С товары разложены по её собственным линейкам, и разложены
плохо: у WELLA 158 позиций свалены в одну кучу «KP Me+/CT», у OLLIN линейки не
проставлены вовсе — все 205 позиций висят под брендом сплошным списком.
Пользователь разложил тот же прайс по человеческим разделам прямо в
Price/Каталог.xlsx («Аммиачный краситель - Koleston», «Оксиданты и осветляющие
средства», «Расходные материалы»). Этот журнал переносит его разбивку на сайт.

ПОЧЕМУ ОТДЕЛЬНОЕ ПОЛЕ, А НЕ ПОДМЕНА `line`. Поле `line` в priceItems.json несёт
на себе ещё две функции помимо группировки:
  * выбор палитры для кружка оттенка — inWellaFamily() смотрит на
    line.startsWith('KP Me+/CT'), ownPaletteFor() тоже разбирает line;
  * привязку бейджей гидов — ключ `БРЕНД :: линейка` в guide-links.json.
Подменить line разделами — значит молча погасить кружки оттенков и потерять
гиды. Поэтому раздел живёт в отдельном поле `section`, а line остаётся как был.

ПОЧЕМУ КЛЮЧ — ИМЯ, А НЕ АРТИКУЛ. Так устроены остальные журналы проекта
(price-categories.json, product-photos.json, price-first-seen.json), и у этого
есть прямая выгода: migrate-renamed-products.py переносит журналы по имени при
переименованиях, поэтому раздел переедет за товаром сам. Имена на сайте берутся
из того же каталога (apply-catalog-names.py), так что они совпадают точно.

Запуск: python scripts/build-price-sections.py [--apply]
Без --apply — предпросмотр. Прогонять после apply-catalog-names.py.
"""
import json
import sys
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / 'Price' / 'Каталог.xlsx'
OUT = ROOT / 'src' / 'data' / 'price-sections.json'
PRICE_ITEMS = ROOT / 'src' / 'data' / 'priceItems.json'


def read_catalog():
    """Бренд — жирный ВЕРХНИЙ регистр 11 кегля, раздел — 11 кегль, товар — 8-й."""
    ws = openpyxl.load_workbook(CATALOG)['Sheet1']
    rows, brand, section = [], None, None
    for r in range(5, ws.max_row + 1):
        cell = ws.cell(row=r, column=2)
        if cell.value is None or not str(cell.value).strip():
            continue
        raw = str(cell.value).strip()
        size = cell.font.size if cell.font and cell.font.size else 0
        bold = bool(cell.font and cell.font.bold)
        is_caps = raw == raw.upper() and any(c.isalpha() for c in raw)
        if bold and is_caps and size >= 11:
            brand, section = raw, None
            continue
        if size >= 11:
            section = raw
            continue
        rows.append({'row': r, 'brand': brand, 'section': section,
                     'name': ' '.join(raw.split())})
    return rows


def main() -> None:
    apply = '--apply' in sys.argv
    rows = read_catalog()
    items = json.loads(PRICE_ITEMS.read_text(encoding='utf-8'))
    site_names = {i['name'] for i in items}

    mapping, no_section, not_on_site = {}, [], []
    for x in rows:
        if not x['section']:
            no_section.append(x)
            continue
        if x['name'] not in site_names:
            not_on_site.append(x)
        mapping[x['name']] = x['section']

    covered = sum(1 for i in items if i['name'] in mapping)
    by_brand = Counter()
    for i in items:
        if i['name'] in mapping:
            by_brand[i.get('brand') or '—'] += 1

    print(f'Каталог: {len(rows)} товаров, разделов у {len(rows) - len(no_section)}')
    print(f'Прайс сайта: {len(items)} позиций, раздел найдётся у {covered}')
    print(f'  без раздела в каталоге: {len(no_section)}')
    print(f'  есть в каталоге, но не на сайте: {len(not_on_site)}')
    print(f'  разных разделов: {len(set(mapping.values()))}')

    print('\nПОКРЫТИЕ ПО БРЕНДАМ:')
    total_by_brand = Counter((i.get('brand') or '—') for i in items)
    for b, total in sorted(total_by_brand.items()):
        got = by_brand.get(b, 0)
        mark = '' if got == total else '   <- не все'
        print(f'  {b:<24} {got:>4} из {total:<4}{mark}')

    if no_section:
        print('\nБЕЗ РАЗДЕЛА — останутся плоским списком под брендом:')
        for x in no_section:
            print(f'  {x["brand"]:<12} {x["name"][:66]}')

    if apply:
        OUT.write_text(
            json.dumps(dict(sorted(mapping.items())), ensure_ascii=False, indent=1) + '\n',
            encoding='utf-8')
        print(f'\nЗаписано: {OUT.relative_to(ROOT)}, {len(mapping)} записей.')
        print('Дальше: xlsx-to-price-items.py, потом сборка.')
    else:
        print('\nЭто предпросмотр, файл не записан. Для записи — с --apply')


if __name__ == '__main__':
    main()
